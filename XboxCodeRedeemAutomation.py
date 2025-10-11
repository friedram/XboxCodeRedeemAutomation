"""
Read.me
What is this? 
When you have many Xbox codes that you want to redeem, using the Xbox PC application is a reliable way to do so.
pyautogui is simply pushing the individual buttons needed in the Xbox app at the right coordinates, copies the
codes out of a word document in a *.docx format. I have a different program that I run to strip out the correct
codes from a larger document and place the codes into the document.
Learn more about pyautogui here https://www.youtube.com/watch?v=lfk_T6VKhTE
ChatGPT o3-mini was used in the creation of this code
You MUST modify the coordinates_first_run and clone some of the coordinates for coordinates_subsequent to make it 
work on your PC.
"""
import time
import pyautogui
import pyperclip
from docx import Document
import os
import re

# Full coordinate sequence for the *first* code redemption
coordinates_first_run = [
    (35, 36),    # "Xbox Face"
    (44, 302),   # "Xbox Settings"
    (725, 659), # "Xbox Redeem"
    (425, 520),   # "Xbox Code" (where we paste the code)
    (524, 952),   # "Xbox NEXT"
    (524, 952),  # "Confirm"
    (524, 952),  # "Close"
    (763, 950)  # "Cancel"
]

# Partial coordinate sequence for *subsequent* redemptions
coordinates_subsequent = [
    (725, 659), # "Xbox Redeem"
    (425, 520),   # "Xbox Code" (where we paste the code)
    (524, 952),   # "Xbox NEXT"
    (524, 952),  # "Confirm"
    (524, 952),  # "Close"
    (763, 950)  # "Cancel"
]

def get_docx_path():
    """
    Prompt the user to input the DOCX file path.
    Continues to prompt until a valid file path is provided.
    """
    example_path = r"D:\\OneDrive\\xbox\\Codes\\Automation\\Code.docx"
    prompt_message = (
        f"Please enter the full path to your DOCX file (e.g., {example_path}): "
    )
    while True:
        docx_path = input(prompt_message).strip()
        if os.path.isfile(docx_path) and docx_path.lower().endswith('.docx'):
            return docx_path
        else:
            print("Invalid file path or file is not a DOCX. Please try again.\n")

def read_all_codes_from_docx(docx_path):
    """
    Searches the DOCX file for codes matching the pattern 'XXXXX-XXXXX-XXXXX-XXXXX-XXXXX'
    using a regex. This function scans both paragraphs and tables.
    Returns a list of code strings.
    """
    document = Document(docx_path)
    codes = []
    pattern = re.compile(r'\b[A-Za-z0-9]{5}(?:-[A-Za-z0-9]{5}){4}\b')
    
    for paragraph in document.paragraphs:
        found = pattern.findall(paragraph.text)
        if found:
            codes.extend(found)
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    found = pattern.findall(paragraph.text)
                    if found:
                        codes.extend(found)
    
    seen = set()
    unique_codes = []
    for code in codes:
        if code not in seen:
            unique_codes.append(code)
            seen.add(code)
    
    return unique_codes

def redeem_code(code, coordinates):
    """
    Goes through the given coordinate list, clicking each coordinate.
    Debugs movement and ensures accuracy.
    """
    for (x, y) in coordinates:
        print(f"Moving to ({x}, {y})")  # Debugging print
        pyautogui.moveTo(x, y, duration=0.8)
        time.sleep(0.5)  # Give time to move
        print(f"Clicking at ({x}, {y})")  # Debugging print
        pyautogui.click()
        
        # Confirm cursor moved
        current_x, current_y = pyautogui.position()
        if (current_x, current_y) != (x, y):
            print(f"Warning: Cursor did not move to expected location. Current position: ({current_x}, {current_y})")
        
        time.sleep(1)  # Short delay to ensure UI response
        
        if (x, y) == (425, 520):
            print("Pasting the code...")
            pyperclip.copy(code)
            pyautogui.hotkey("ctrl", "v")
        
        time.sleep(7)

def main():
    docx_path = get_docx_path()
    
    codes = read_all_codes_from_docx(docx_path)
    if not codes:
        print("No valid codes found in the DOCX file.")
        return
    
    print(f"Found {len(codes)} code(s):")
    for c in codes:
        print(" -", c)
    
    print(f"\nThe total codes found in the file are: {len(codes)}")
    
    print("\nPlease make sure the Xbox app is open and visible. You have 5 seconds...")
    time.sleep(5)
    
    print(f"\nRedeeming first code (attempt 1): {codes[0]}")
    redeem_code(codes[0], coordinates_first_run)
    print(f"\nRedeeming first code (attempt 2): {codes[0]}")
    redeem_code(codes[0], coordinates_first_run)
    
    for code in codes[1:]:
        print(f"\nRedeeming next code: {code}")
        redeem_code(code, coordinates_subsequent)
        time.sleep(5)

    print("\nAll codes processed.")

if __name__ == "__main__":
    main()
